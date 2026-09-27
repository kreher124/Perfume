"""Helpers for matching a perfume sheet to Fragrantica. Run from the repo root.

  python3 scripts/sheet_tools.py unmatched [sheet.xlsx]
      Lists every perfume in the sheet that has no entry in matches.json yet,
      as "tab | house | name". These are the ones still to look up.

  python3 scripts/sheet_tools.py corrected [sheet.xlsx] [out.xlsx]
      Writes a copy of the sheet with house and perfume names replaced by the
      corrected ones from matches.json and HOUSE_FIXES in app.js. Scores, notes
      and cell colors are left alone. The owner imports it into Google Sheets
      with File > Import > Upload > Replace spreadsheet.

The sheet defaults to data/FRAGRANCES.xlsx, the hourly backup of the live sheet.
Needs openpyxl (pip install openpyxl).
"""
import json
import re
import sys
import unicodedata

import openpyxl


def norm(s):
    s = unicodedata.normalize("NFD", str(s if s is not None else ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def clean(v):
    # Google exports a name like 801 as the number 801.0.
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def house_fixes():
    src = open("app.js", encoding="utf-8").read()
    body = re.search(r"HOUSE_FIXES = \{(.*?)\};", src, re.S).group(1)
    pairs = re.findall(r'^\s*"((?:[^"\\]|\\.)*)": "((?:[^"\\]|\\.)*)",$', body, re.M)
    return {norm(json.loads(f'"{k}"')): json.loads(f'"{v}"') for k, v in pairs}


def load_matches(fixes):
    """matches.json, also filed under each entry's corrected names, as app.js does."""
    matches = json.load(open("matches.json", encoding="utf-8"))
    for key, m in list(matches.items()):
        h, n = key.split("|", 1)
        matches.setdefault(norm(m.get("h") or fixes.get(h) or h) + "|" + norm(m.get("n") or n), m)
    return matches


def rows(wb):
    """Yields (tab, house cell, name cell) for every perfume row, like app.js reads them."""
    for ws in wb.worksheets:
        hdr = [norm(c.value) for c in ws[1]]
        hi = hdr.index("house") if "house" in hdr else 0
        ni = hdr.index("name") if "name" in hdr else 1
        for r in ws.iter_rows(min_row=2):
            h, n = clean(r[hi].value), clean(r[ni].value)
            if not h and not n:
                continue
            if not n and h.startswith("("):  # a note, not a perfume
                continue
            yield ws.title, r[hi], r[ni]


def unmatched(path="data/FRAGRANCES.xlsx"):
    matches = load_matches(house_fixes())
    seen = set()
    for tab, hc, nc in rows(openpyxl.load_workbook(path)):
        h, n = clean(hc.value), clean(nc.value)
        k = norm(h) + "|" + norm(n)
        if k not in matches and k not in seen:
            seen.add(k)
            print(f"{tab} | {h} | {n}")
    print(f"{len(seen)} not in matches.json", file=sys.stderr)


def corrected(path="data/FRAGRANCES.xlsx", out="FRAGRANCES (corrected).xlsx"):
    fixes = house_fixes()
    matches = load_matches(fixes)
    wb = openpyxl.load_workbook(path)
    changed = 0
    for _, hc, nc in rows(wb):
        h, n = clean(hc.value), clean(nc.value)
        m = matches.get(norm(h) + "|" + norm(n)) or {}
        dh = m.get("h") or fixes.get(norm(h)) or h
        dn = m.get("n") or n
        # keep the owner's trailing notes, like "(rose)" or "- 10 ml"
        suffix = re.search(r"(\s*\((?!or )[^)]*\)|\s+-\s+\d+\s*ml)$", n)
        if n and suffix and not dn.endswith(suffix.group(1)):
            dn += suffix.group(1)
        if (dh, dn) != (h, n):
            changed += 1
        hc.value = dh  # also trims stray spaces
        if n:
            nc.value = dn
    wb.save(out)
    print(f"{changed} rows corrected, saved to {out}", file=sys.stderr)


if __name__ == "__main__":
    cmd, *args = sys.argv[1:] or ["help"]
    {"unmatched": unmatched, "corrected": corrected}.get(cmd, lambda *a: print(__doc__))(*args)
